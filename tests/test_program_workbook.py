import pandas as pd
from openpyxl import load_workbook

from src.program_workbook import (
    PROGRAM_CIP_MAP_COLUMNS,
    PROGRAM_INVENTORY_COLUMNS,
    build_program_onepager_data,
    load_program_cip_map,
    load_program_inventory,
    write_program_workbook,
)


def _summary():
    return pd.DataFrame(
        {
            "cip": ["30.9999", "45.0901"],
            "program_name": [
                "Multi-/Interdisciplinary Studies, Other",
                "International Relations and Affairs",
            ],
            "cip_title": [
                "Multi-/Interdisciplinary Studies, Other",
                "International Relations and Affairs",
            ],
            "awards": [121, 386],
            "occupation_count": [0, 3],
            "soc_match": [False, True],
            "wtd_median_wage_national": [pd.NA, 140748.0],
        }
    )


def _detail():
    return pd.DataFrame(
        {
            "cip": ["30.9999", "45.0901"],
            "program_name": [
                "Multi-/Interdisciplinary Studies, Other",
                "International Relations and Affairs",
            ],
            "cip_title": [
                "Multi-/Interdisciplinary Studies, Other",
                "International Relations and Affairs",
            ],
            "soc": [pd.NA, "11-9199"],
            "soc_title": [pd.NA, "Managers, All Other"],
            "soc_match": [False, True],
            "do_not_display_on_one_pager": [False, True],
            "one_pager_display_note": [pd.NA, "residual catch-all occupation"],
        }
    )


def test_blank_local_inputs_generate_unpublishable_cip_placeholders():
    empty_inventory = pd.DataFrame(columns=PROGRAM_INVENTORY_COLUMNS)
    empty_cip_map = pd.DataFrame(columns=PROGRAM_CIP_MAP_COLUMNS)
    program_summary, program_occ, inventory, cip_map = build_program_onepager_data(
        _summary(), _detail(), empty_inventory, empty_cip_map
    )

    assert len(program_summary) == 2
    row = program_summary[program_summary["cip"] == "30.9999"].iloc[0]
    assert row["program_id"] == "cip_30_9999"
    assert row["program_name"] == ""
    assert row["cip_title"] == "Multi-/Interdisciplinary Studies, Other"
    assert row["cip_mapping_type"] == "provisional_cip_row"
    assert row["one_pager_ready"] == False
    assert pd.isna(row["program_awards_conferred_2022_23"])
    assert len(inventory) == 2
    assert len(cip_map) == 2
    assert "program_id" in program_occ.columns


def test_verified_one_to_one_program_gets_program_key_and_awards():
    inventory = pd.DataFrame(
        {
            **{c: [""] for c in PROGRAM_INVENTORY_COLUMNS},
            "program_id": ["ir_ba"],
            "program_name": ["International Studies"],
            "program_name_verified": ["TRUE"],
            "degree_level": ["Bachelor's"],
        }
    )
    cip_map = pd.DataFrame(
        {
            **{c: [""] for c in PROGRAM_CIP_MAP_COLUMNS},
            "program_id": ["ir_ba"],
            "cip": ["45.0901"],
            "cip_mapping_type": ["one_program_to_one_cip"],
            "cip_mapping_verified": ["TRUE"],
        }
    )

    program_summary, program_occ, _, _ = build_program_onepager_data(
        _summary(), _detail(), inventory, cip_map
    )

    row = program_summary.iloc[0]
    assert row["program_id"] == "ir_ba"
    assert row["program_name"] == "International Studies"
    assert row["cip_title"] == "International Relations and Affairs"
    assert row["one_pager_ready"] == True
    assert row["program_awards_conferred_2022_23"] == 386
    assert program_occ.iloc[0]["program_id"] == "ir_ba"
    assert program_occ.iloc[0]["do_not_display_on_one_pager"] == True


def test_labor_market_proxy_borrows_occupations_but_keeps_reported_cip_awards():
    inventory = pd.DataFrame(
        {
            **{c: [""] for c in PROGRAM_INVENTORY_COLUMNS},
            "program_id": ["cleg_ba"],
            "program_name": ["CLEG"],
            "program_name_verified": ["TRUE"],
            "degree_level": ["Bachelor's"],
        }
    )
    cip_map = pd.DataFrame(
        {
            **{c: [""] for c in PROGRAM_CIP_MAP_COLUMNS},
            "program_id": ["cleg_ba"],
            "cip": ["30.9999"],
            "labor_market_cip": ["45.0901"],
            "cip_mapping_type": ["multiple_programs_share_cip"],
            "cip_mapping_verified": ["TRUE"],
        }
    )

    program_summary, program_occ, _, _ = build_program_onepager_data(
        _summary(), _detail(), inventory, cip_map
    )

    row = program_summary.iloc[0]
    assert row["cip"] == "30.9999"
    assert row["labor_market_cip"] == "45.0901"
    assert row["cip_awards_conferred_2022_23"] == 121  # awards from the reported CIP
    assert row["occupation_count"] == 3  # occupations from the proxy CIP
    assert row["wtd_median_wage_national"] == 140748.0  # wage from the proxy CIP
    assert row["cip_labor_market_match"] == True
    assert row["one_pager_ready"] == True
    assert row["labor_market_data_level"] == "CIP (proxy)"
    assert "proxy" in str(row["labor_market_proxy_note"]).lower()
    assert program_occ.iloc[0]["soc"] == "11-9199"  # proxy CIP's occupation


def test_verified_program_without_cip_labor_market_match_is_blocked():
    inventory = pd.DataFrame(
        {
            **{c: [""] for c in PROGRAM_INVENTORY_COLUMNS},
            "program_id": ["dance_ba"],
            "program_name": ["Dance"],
            "program_name_verified": ["TRUE"],
            "degree_level": ["Bachelor's"],
        }
    )
    cip_map = pd.DataFrame(
        {
            **{c: [""] for c in PROGRAM_CIP_MAP_COLUMNS},
            "program_id": ["dance_ba"],
            "cip": ["50.0301"],
            "cip_mapping_type": ["one_program_to_one_cip"],
            "cip_mapping_verified": ["TRUE"],
        }
    )

    program_summary, program_occ, _, _ = build_program_onepager_data(
        _summary(), _detail(), inventory, cip_map
    )

    row = program_summary.iloc[0]
    assert row["program_id"] == "dance_ba"
    assert row["cip_labor_market_match"] == False
    assert row["one_pager_ready"] == False
    assert "No CIP labor-market summary match" in row["one_pager_blocker"]
    assert program_occ.iloc[0]["cip_labor_market_match"] == False


def test_lookup_loaders_validate_columns(tmp_path):
    inv = tmp_path / "program_inventory.csv"
    pd.DataFrame(columns=PROGRAM_INVENTORY_COLUMNS).to_csv(inv, index=False)
    # An inventory without the optional brief_note column still loads, reading it blank.
    loaded = load_program_inventory(inv)
    assert list(loaded.columns) == PROGRAM_INVENTORY_COLUMNS + ["brief_note"]

    cip_map = tmp_path / "program_cip_map.csv"
    pd.DataFrame(columns=PROGRAM_CIP_MAP_COLUMNS).to_csv(cip_map, index=False)
    assert list(load_program_cip_map(cip_map).columns) == PROGRAM_CIP_MAP_COLUMNS


def test_program_workbook_writes_expected_sheets(tmp_path):
    program_summary, program_occ, inventory, cip_map = build_program_onepager_data(
        _summary(), _detail()
    )
    out = tmp_path / "programs.xlsx"
    write_program_workbook(
        program_summary,
        program_occ,
        inventory,
        cip_map,
        [("Purpose", "test")],
        out,
    )
    wb = load_workbook(out)
    assert wb.sheetnames == [
        "Program One-Pager Data",
        "Program Occupations",
        "Program Inventory",
        "Program-CIP Map",
        "Methodology",
    ]
